from pptx import Presentation

from spidermapp.core import pptx_report
from spidermapp.core.models import CrawlResult, IssueCategory, IssueSeverity, PageResult


def _result():
    result = CrawlResult(seed_url="https://example.com/")
    page = PageResult(url="https://example.com/", status_code=200)
    page.add_issue(IssueCategory.TITLES, IssueSeverity.CRITICAL, "title_missing", "Falta el title.")
    result.pages.append(page)
    error_page = PageResult(url="https://example.com/roto", status_code=404)
    result.pages.append(error_page)
    return result


def test_default_pptx_filename_uses_bare_domain():
    assert pptx_report.default_pptx_filename("https://www.cruzverde.com.co/tienda") == "presentación_status_cruzverde.com.co.pptx"


def test_default_pptx_filename_falls_back_when_url_unparseable():
    assert pptx_report.default_pptx_filename("") == "presentación_status_sitio.pptx"


def test_generate_pptx_report_writes_valid_file(tmp_path):
    out = tmp_path / "presentacion.pptx"
    path = pptx_report.generate_pptx_report(_result(), out)
    assert path == out
    assert out.exists()

    prs = Presentation(str(out))
    assert len(prs.slides) == 5  # título, resumen, categorías, oportunidades, cierre


def test_generate_pptx_report_handles_empty_crawl(tmp_path):
    out = tmp_path / "empty.pptx"
    result = CrawlResult(seed_url="https://example.com/")
    path = pptx_report.generate_pptx_report(result, out)
    assert out.exists()
    prs = Presentation(str(out))
    assert len(prs.slides) == 5


def test_generate_pptx_report_slide_dimensions_are_widescreen(tmp_path):
    out = tmp_path / "widescreen.pptx"
    pptx_report.generate_pptx_report(_result(), out)
    prs = Presentation(str(out))
    assert prs.slide_width == pptx_report.SLIDE_W
    assert prs.slide_height == pptx_report.SLIDE_H
