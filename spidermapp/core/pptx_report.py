from __future__ import annotations

from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

from spidermapp.core import recommendations, stats, url_utils
from spidermapp.core.models import CrawlResult, Issue

# Sober, modern palette — matches the app's own theme (indigo primary,
# neutral grays) rather than a generic template look.
INK = RGBColor(0x11, 0x18, 0x27)
MUTED = RGBColor(0x6B, 0x72, 0x80)
SUBTLE = RGBColor(0x9C, 0xA3, 0xAF)
PRIMARY = RGBColor(0x4F, 0x46, 0xE5)
PRIMARY_DARK = RGBColor(0x37, 0x30, 0xA3)
PRIMARY_SOFT = RGBColor(0xEE, 0xF2, 0xFF)
BG_LIGHT = RGBColor(0xF9, 0xFA, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOOD = RGBColor(0x05, 0x96, 0x69)
GOOD_SOFT = RGBColor(0xEC, 0xFD, 0xF5)
WARNING = RGBColor(0xD9, 0x77, 0x06)
WARNING_SOFT = RGBColor(0xFF, 0xFB, 0xEB)
CRITICAL = RGBColor(0xDC, 0x26, 0x26)
CRITICAL_SOFT = RGBColor(0xFE, 0xF2, 0xF2)
BORDER = RGBColor(0xE5, 0xE7, 0xEB)

FONT = "Helvetica Neue"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

OPPORTUNITIES_PER_SLIDE = 4
MAX_OPPORTUNITIES = 12
MAX_ERROR_ROWS = 10
MAX_DUPLICATE_GROUPS = 3


def default_pptx_filename(seed_url: str) -> str:
    domain = url_utils.strip_www(url_utils.registrable_domain(seed_url)) or "sitio"
    return f"presentación_status_{domain}.pptx"


def _score_color(score: int) -> RGBColor:
    if score < 50:
        return CRITICAL
    if score < 80:
        return WARNING
    return GOOD


def _blank_slide(prs: Presentation, bg: RGBColor = WHITE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    rect.fill.solid()
    rect.fill.fore_color.rgb = bg
    rect.line.fill.background()
    rect.shadow.inherit = False
    # Push the background rect behind everything added afterward.
    spTree = slide.shapes._spTree
    spTree.remove(rect._element)
    spTree.insert(2, rect._element)
    return slide


def _textbox(slide, left, top, width, height, text, size, color, bold=False, align=PP_ALIGN.LEFT, font=FONT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font
    return box


def _bullets(slide, left, top, width, height, items, size, color, spacing=6, font=FONT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = f"•  {item}"
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = font
    return box


def _slide_header(slide, title: str, subtitle: str = "") -> None:
    _textbox(slide, Inches(0.6), Inches(0.45), Inches(11.5), Inches(0.6), title, 26, INK, bold=True)
    if subtitle:
        _textbox(slide, Inches(0.6), Inches(1.05), Inches(11.8), Inches(0.4), subtitle, 12, MUTED)


def _footer(slide, prs, seed_url: str, page_label: str) -> None:
    _textbox(
        slide, Inches(0.5), prs.slide_height - Inches(0.45), Inches(7), Inches(0.3),
        url_utils.strip_www(url_utils.registrable_domain(seed_url)), 9, SUBTLE,
    )
    _textbox(
        slide, prs.slide_width - Inches(3.5), prs.slide_height - Inches(0.45), Inches(3), Inches(0.3),
        page_label, 9, SUBTLE, align=PP_ALIGN.RIGHT,
    )


def _kpi_card(slide, left, top, width, height, value: str, label: str, value_color: RGBColor) -> None:
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.adjustments[0] = 0.08
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = BORDER
    card.line.width = Pt(0.75)
    card.shadow.inherit = False

    _textbox(slide, left, top + Inches(0.18), width, Inches(0.7), value, 30, value_color, bold=True, align=PP_ALIGN.CENTER)
    _textbox(slide, left, top + height - Inches(0.55), width, Inches(0.4), label, 10.5, MUTED, align=PP_ALIGN.CENTER)


def _pill(slide, left, top, ok: bool, ok_text: str, bad_text: str) -> Emu:
    """Draws a status pill and returns its actual width, so callers can lay
    out the next one right after it."""
    text = f"✓  {ok_text}" if ok else f"✕  {bad_text}"
    color = GOOD if ok else CRITICAL
    bg = GOOD_SOFT if ok else CRITICAL_SOFT
    width = Inches(0.22 + len(text) * 0.078)
    height = Inches(0.42)
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    pill.adjustments[0] = 0.5
    pill.fill.solid()
    pill.fill.fore_color.rgb = bg
    pill.line.fill.background()
    pill.shadow.inherit = False
    tf = pill.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = color
    run.font.name = FONT
    return width


def generate_pptx_report(result: CrawlResult, path: str | Path, priority_urls: set[str] | None = None) -> Path:
    priority_urls = priority_urls or set()
    score = stats.health_score(result.pages, priority_urls)
    budget = stats.compute_crawl_budget(result.pages)
    breakdown = stats.category_breakdown(result.pages)
    opportunities = stats.find_opportunities(result.pages, limit=MAX_OPPORTUNITIES)
    errors = stats.error_pages(result.pages)
    infra = stats.summarize_site_infrastructure(result.sitemap_urls, result.site_issues)
    domain = url_utils.strip_www(url_utils.registrable_domain(result.seed_url)) or result.seed_url

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    _build_title_slide(prs, domain, result, budget)
    _build_summary_slide(prs, result, budget, score)
    _build_infra_slide(prs, result, infra)
    _build_categories_slide(prs, result, breakdown)
    _build_errors_slide(prs, result, errors)
    _build_opportunities_slides(prs, result, opportunities)
    if result.duplicate_content_groups:
        _build_duplicates_slide(prs, result)
    _build_closing_slide(prs, result, budget, breakdown, opportunities)

    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    return output_path


def _build_title_slide(prs: Presentation, domain: str, result: CrawlResult, budget: stats.CrawlBudget) -> None:
    slide = _blank_slide(prs, bg=INK)

    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), prs.slide_height)
    accent.fill.solid()
    accent.fill.fore_color.rgb = PRIMARY
    accent.line.fill.background()
    accent.shadow.inherit = False

    _textbox(slide, Inches(1), Inches(2.3), Inches(11), Inches(0.5), "AUDITORÍA SEO", 16, RGBColor(0xA5, 0xB4, 0xFC), bold=True)
    _textbox(slide, Inches(1), Inches(2.8), Inches(11), Inches(1.1), domain, 44, WHITE, bold=True)
    _textbox(slide, Inches(1), Inches(3.8), Inches(11), Inches(0.5), "Resumen ejecutivo de status del sitio", 18, RGBColor(0xC7, 0xCC, 0xD8))
    _textbox(
        slide, Inches(1), Inches(4.35), Inches(11), Inches(0.4),
        f"{result.seed_url}   ·   {datetime.now().strftime('%d/%m/%Y')}", 12, SUBTLE,
    )

    duration = ""
    if result.finished_at:
        seconds = max(0, result.finished_at - result.started_at)
        duration = f"   ·   {int(seconds // 60)} min {int(seconds % 60)} s" if seconds >= 60 else f"   ·   {int(seconds)} s"
    _textbox(
        slide, Inches(1), Inches(4.75), Inches(11), Inches(0.4),
        f"{budget.total_pages} URLs analizadas{duration}", 12, SUBTLE,
    )

    _textbox(slide, Inches(1), Inches(6.6), Inches(4), Inches(0.4), "PREPARADO CON SPIDERMAPP", 10, RGBColor(0x6B, 0x72, 0x80), bold=True)


def _build_summary_slide(prs: Presentation, result: CrawlResult, budget: stats.CrawlBudget, score: int) -> None:
    slide = _blank_slide(prs, bg=BG_LIGHT)
    _slide_header(slide, "Resumen ejecutivo", "Estado general del sitio a la fecha de este análisis.")

    # Health score hero
    ring = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.7), Inches(1.7), Inches(2.1), Inches(2.1))
    ring.fill.solid()
    ring.fill.fore_color.rgb = WHITE
    ring.line.color.rgb = _score_color(score)
    ring.line.width = Pt(4)
    ring.shadow.inherit = False
    tf = ring.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = str(score)
    run.font.size = Pt(40)
    run.font.bold = True
    run.font.color.rgb = _score_color(score)
    run.font.name = FONT

    _textbox(slide, Inches(0.55), Inches(3.85), Inches(2.4), Inches(0.4), "Puntaje de salud", 12, MUTED, align=PP_ALIGN.CENTER)

    ok_pct = round(100 * budget.ok_pages / budget.total_pages) if budget.total_pages else 0
    _textbox(slide, Inches(0.55), Inches(4.25), Inches(2.4), Inches(0.35), f"{ok_pct}% de las páginas sin errores", 10.5, MUTED, align=PP_ALIGN.CENTER)

    # KPI cards
    cards = [
        (str(budget.total_pages), "URLs rastreadas", INK),
        (str(budget.ok_pages), "Páginas OK", GOOD),
        (str(budget.error_pages), "Con errores", CRITICAL),
        (str(budget.duplicate_pages), "Contenido duplicado", WARNING),
    ]
    card_w, card_h, gap = Inches(2.55), Inches(1.5), Inches(0.25)
    start_x = Inches(3.4)
    for i, (value, label, color) in enumerate(cards):
        x = start_x + i * (card_w + gap)
        _kpi_card(slide, x, Inches(1.7), card_w, card_h, value, label, color)

    cards2 = [
        (str(budget.blocked_by_robots), "Bloqueadas por robots.txt", MUTED),
        (str(budget.orphan_pages), "Páginas huérfanas", WARNING),
    ]
    for i, (value, label, color) in enumerate(cards2):
        x = start_x + i * (card_w + gap)
        _kpi_card(slide, x, Inches(3.4), card_w, card_h, value, label, color)

    note = (
        "El puntaje de salud combina la severidad y cantidad de hallazgos técnicos detectados; "
        "100 indica ausencia de problemas relevantes."
    )
    _textbox(slide, Inches(3.4), Inches(5.15), Inches(9.2), Inches(1.2), note, 11, MUTED)

    _footer(slide, prs, result.seed_url, "01  ·  Resumen ejecutivo")


def _build_infra_slide(prs: Presentation, result: CrawlResult, infra: stats.SiteInfrastructure) -> None:
    slide = _blank_slide(prs, bg=WHITE)
    _slide_header(slide, "Salud técnica: robots.txt y sitemap", "Señales que determinan si los motores de búsqueda pueden rastrear e indexar el sitio.")

    y = Inches(1.9)
    x = Inches(0.6)
    x += _pill(slide, x, y, infra.robots_found, "robots.txt encontrado", "Sin robots.txt") + Inches(0.18)
    if infra.robots_found:
        x += _pill(slide, x, y, infra.robots_allows_crawling, "Permite el rastreo", "Bloquea todo el sitio") + Inches(0.18)
        _pill(slide, x, y, infra.robots_declares_sitemap, "Declara el sitemap", "No declara el sitemap")

    y2 = Inches(2.55)
    sitemap_label = f"Sitemap encontrado — {infra.sitemap_url_count:,} URLs".replace(",", ".")
    _pill(slide, Inches(0.6), y2, infra.sitemap_found, sitemap_label, "Sin sitemap o está vacío")

    other_top = Inches(3.5)
    if infra.other_findings:
        _textbox(slide, Inches(0.6), other_top, Inches(11), Inches(0.4), "Otros hallazgos técnicos a nivel de sitio", 14, INK, bold=True)
        row_h = Inches(0.5)
        for i, issue in enumerate(infra.other_findings[:8]):
            rec = recommendations.get_recommendation(issue)
            y3 = other_top + Inches(0.55) + i * row_h
            color = CRITICAL if issue.severity.value == "critical" else (WARNING if issue.severity.value == "warning" else PRIMARY)
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.65), y3 + Inches(0.08), Inches(0.12), Inches(0.12))
            dot.fill.solid()
            dot.fill.fore_color.rgb = color
            dot.line.fill.background()
            dot.shadow.inherit = False
            _textbox(slide, Inches(0.95), y3, Inches(11), Inches(0.4), rec.title, 12.5, INK)
    else:
        _textbox(slide, Inches(0.6), other_top, Inches(11), Inches(0.4), "Sin otros hallazgos técnicos a nivel de sitio. ✓", 13, GOOD, bold=True)

    _footer(slide, prs, result.seed_url, "02  ·  Salud técnica")


def _build_categories_slide(prs: Presentation, result: CrawlResult, breakdown: list[stats.CategoryBreakdown]) -> None:
    slide = _blank_slide(prs, bg=BG_LIGHT)
    _slide_header(slide, "Hallazgos por categoría", "Distribución de issues técnicos detectados, por área de SEO.")

    top = Inches(1.75)
    row_h = Inches(0.55)
    max_total = max((b.total for b in breakdown), default=1) or 1
    bar_area_x = Inches(4.6)
    bar_area_w = Inches(6.6)

    if not breakdown:
        _textbox(slide, Inches(0.6), top, Inches(10), Inches(0.5), "Sin issues detectados en este análisis.", 13, GOOD, bold=True)
    else:
        total_issues = sum(b.total for b in breakdown)
        for i, bucket in enumerate(breakdown[:9]):
            y = top + i * row_h
            label = bucket.category.value.replace("_", " ").title()
            _textbox(slide, Inches(0.6), y, Inches(3.8), Inches(0.4), label, 11.5, INK)

            track = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bar_area_x, y + Emu(60000), bar_area_w, Inches(0.28))
            track.adjustments[0] = 0.5
            track.fill.solid()
            track.fill.fore_color.rgb = BORDER
            track.line.fill.background()
            track.shadow.inherit = False

            frac = bucket.total / max_total
            fill_w = max(Emu(int(bar_area_w * frac)), Inches(0.05))
            color = CRITICAL if bucket.critical >= bucket.warning and bucket.critical >= bucket.info else (
                WARNING if bucket.warning >= bucket.info else PRIMARY
            )
            fill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bar_area_x, y + Emu(60000), fill_w, Inches(0.28))
            fill.adjustments[0] = 0.5
            fill.fill.solid()
            fill.fill.fore_color.rgb = color
            fill.line.fill.background()
            fill.shadow.inherit = False

            pct = round(100 * bucket.total / total_issues) if total_issues else 0
            _textbox(slide, bar_area_x + bar_area_w + Inches(0.15), y, Inches(1.7), Inches(0.4), f"{bucket.total}  ({pct}%)", 11.5, MUTED)

    _footer(slide, prs, result.seed_url, "03  ·  Hallazgos por categoría")


def _build_errors_slide(prs: Presentation, result: CrawlResult, errors: list) -> None:
    slide = _blank_slide(prs, bg=WHITE)
    _slide_header(slide, f"Páginas con error ({len(errors)} en total)", "Códigos de respuesta 4xx/5xx o solicitudes sin respuesta del servidor.")

    if not errors:
        _textbox(slide, Inches(0.6), Inches(1.9), Inches(10), Inches(0.5), "Sin páginas con errores de respuesta. ✓", 14, GOOD, bold=True)
    else:
        top = Inches(1.85)
        row_h = Inches(0.46)
        header_h = Inches(0.4)
        table_left = Inches(0.6)
        status_w = Inches(1.3)
        url_w = Inches(10.9)

        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, table_left, top, status_w + url_w, header_h)
        header.fill.solid()
        header.fill.fore_color.rgb = RGBColor(0xF3, 0xF4, 0xF6)
        header.line.fill.background()
        header.shadow.inherit = False
        _textbox(slide, table_left + Inches(0.15), top + Inches(0.07), status_w, Inches(0.3), "ESTADO", 10, MUTED, bold=True)
        _textbox(slide, table_left + status_w + Inches(0.15), top + Inches(0.07), url_w, Inches(0.3), "URL", 10, MUTED, bold=True)

        for i, page in enumerate(errors[:MAX_ERROR_ROWS]):
            y = top + header_h + i * row_h
            label = page.error[:16] if page.error else (str(page.status_code) if page.status_code is not None else "Sin respuesta")
            _textbox(slide, table_left + Inches(0.15), y + Inches(0.08), status_w, Inches(0.32), label, 11, CRITICAL, bold=True)
            url_text = page.url if len(page.url) <= 100 else page.url[:97] + "…"
            _textbox(slide, table_left + status_w + Inches(0.15), y + Inches(0.08), url_w, Inches(0.32), url_text, 10.5, INK)
            if i < min(len(errors), MAX_ERROR_ROWS) - 1:
                rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, table_left, y + row_h, status_w + url_w, Pt(0.75))
                rule.fill.solid()
                rule.fill.fore_color.rgb = BORDER
                rule.line.fill.background()
                rule.shadow.inherit = False

        if len(errors) > MAX_ERROR_ROWS:
            _textbox(
                slide, table_left, top + header_h + MAX_ERROR_ROWS * row_h + Inches(0.15), Inches(11), Inches(0.35),
                f"... y {len(errors) - MAX_ERROR_ROWS} páginas con error adicionales (ver reporte PDF para el listado completo).",
                10.5, MUTED,
            )

    _footer(slide, prs, result.seed_url, "04  ·  Errores")


def _build_opportunities_slides(prs: Presentation, result: CrawlResult, opportunities: list[stats.IssueOpportunity]) -> None:
    if not opportunities:
        slide = _blank_slide(prs, bg=BG_LIGHT)
        _slide_header(slide, "Principales oportunidades", "Acciones priorizadas con mayor impacto potencial para el sitio.")
        _textbox(slide, Inches(0.6), Inches(1.8), Inches(10), Inches(0.5), "Sin oportunidades relevantes detectadas — el sitio está en buena forma.", 13, GOOD, bold=True)
        _footer(slide, prs, result.seed_url, "05  ·  Oportunidades")
        return

    chunks = [opportunities[i : i + OPPORTUNITIES_PER_SLIDE] for i in range(0, len(opportunities), OPPORTUNITIES_PER_SLIDE)]
    for part, chunk in enumerate(chunks, start=1):
        slide = _blank_slide(prs, bg=BG_LIGHT)
        suffix = f" ({part}/{len(chunks)})" if len(chunks) > 1 else ""
        _slide_header(slide, f"Principales oportunidades{suffix}", "Acciones priorizadas con mayor impacto potencial para el sitio.")

        top = Inches(1.75)
        card_h = Inches(1.2)
        gap = Inches(0.18)
        for i, opp in enumerate(chunk):
            rec = recommendations.get_recommendation(Issue(opp.category, opp.severity, opp.code, ""))
            y = top + i * (card_h + gap)
            color = CRITICAL if opp.severity.value == "critical" else (WARNING if opp.severity.value == "warning" else PRIMARY)

            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), y, Inches(12.1), card_h)
            card.adjustments[0] = 0.08
            card.fill.solid()
            card.fill.fore_color.rgb = WHITE
            card.line.color.rgb = BORDER
            card.line.width = Pt(0.75)
            card.shadow.inherit = False

            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), y, Inches(0.09), card_h)
            bar.fill.solid()
            bar.fill.fore_color.rgb = color
            bar.line.fill.background()
            bar.shadow.inherit = False

            pages_word = "página" if opp.affected_pages == 1 else "páginas"
            _textbox(slide, Inches(0.9), y + Inches(0.12), Inches(9.6), Inches(0.35), rec.title, 13.5, INK, bold=True)
            if rec.why:
                _textbox(slide, Inches(0.9), y + Inches(0.47), Inches(9.6), Inches(0.32), f"Por qué importa: {rec.why}", 10, MUTED)
            if rec.fix_steps:
                _textbox(slide, Inches(0.9), y + Inches(0.79), Inches(9.6), Inches(0.32), f"Solución: {rec.fix_steps[0]}", 10, color)
            _textbox(
                slide, Inches(10.7), y + Inches(0.42), Inches(1.9), Inches(0.4),
                f"{opp.affected_pages} {pages_word}", 12, color, bold=True, align=PP_ALIGN.RIGHT,
            )

        _footer(slide, prs, result.seed_url, f"05  ·  Oportunidades{suffix}")


def _build_duplicates_slide(prs: Presentation, result: CrawlResult) -> None:
    slide = _blank_slide(prs, bg=WHITE)
    groups = result.duplicate_content_groups
    _slide_header(
        slide,
        f"Contenido duplicado ({len(groups)} grupos)",
        "Conjuntos de páginas con contenido prácticamente idéntico — compiten entre sí por las mismas búsquedas.",
    )

    top = Inches(1.85)
    row_h_group = Inches(1.35)
    for i, (content_hash, urls) in enumerate(list(groups.items())[:MAX_DUPLICATE_GROUPS]):
        y = top + i * row_h_group
        _textbox(slide, Inches(0.6), y, Inches(11.5), Inches(0.35), f"Grupo {i + 1}  —  {len(urls)} páginas afectadas", 13, INK, bold=True)
        shown = urls[:4]
        for j, url in enumerate(shown):
            url_text = url if len(url) <= 100 else url[:97] + "…"
            _textbox(slide, Inches(0.9), y + Inches(0.4) + j * Inches(0.26), Inches(11), Inches(0.28), f"•  {url_text}", 10.5, MUTED)
        if len(urls) > len(shown):
            _textbox(slide, Inches(0.9), y + Inches(0.4) + len(shown) * Inches(0.26), Inches(11), Inches(0.28), f"...y {len(urls) - len(shown)} más.", 10, SUBTLE)

    if len(groups) > MAX_DUPLICATE_GROUPS:
        _textbox(
            slide, Inches(0.6), top + MAX_DUPLICATE_GROUPS * row_h_group, Inches(11), Inches(0.35),
            f"... y {len(groups) - MAX_DUPLICATE_GROUPS} grupos de contenido duplicado adicionales.", 10.5, MUTED,
        )

    _footer(slide, prs, result.seed_url, "06  ·  Contenido duplicado")


def _build_closing_slide(
    prs: Presentation,
    result: CrawlResult,
    budget: stats.CrawlBudget,
    breakdown: list[stats.CategoryBreakdown],
    opportunities: list[stats.IssueOpportunity],
) -> None:
    slide = _blank_slide(prs, bg=INK)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), prs.slide_height)
    accent.fill.solid()
    accent.fill.fore_color.rgb = PRIMARY
    accent.line.fill.background()
    accent.shadow.inherit = False

    _textbox(slide, Inches(1), Inches(1.4), Inches(11), Inches(0.7), "Próximos pasos", 30, WHITE, bold=True)

    steps: list[str] = []
    critical_opps = [o for o in opportunities if o.severity.value == "critical"]
    if critical_opps:
        top_opp = recommendations.get_recommendation(Issue(critical_opps[0].category, critical_opps[0].severity, critical_opps[0].code, ""))
        steps.append(f"Priorizar: {top_opp.title.lower()} ({critical_opps[0].affected_pages} páginas afectadas).")
    if breakdown:
        top_category = breakdown[0].category.value.replace("_", " ").title()
        steps.append(f"La categoría con más hallazgos es \"{top_category}\" — enfocar el primer sprint de correcciones ahí.")
    if budget.error_pages:
        steps.append(f"Resolver las {budget.error_pages} páginas con error para recuperar presupuesto de rastreo.")
    steps.append("Reevaluar el sitio después de aplicar las correcciones para medir el impacto.")
    steps.append("Programar auditorías periódicas para detectar regresiones a tiempo.")

    _bullets(slide, Inches(1), Inches(2.3), Inches(11), Inches(3.8), steps, 14.5, RGBColor(0xC7, 0xCC, 0xD8), spacing=12)
    _textbox(slide, Inches(1), Inches(6.9), Inches(6), Inches(0.4), "Generado con Spidermapp", 10, SUBTLE)
